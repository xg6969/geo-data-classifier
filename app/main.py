import os
import shutil
import zipfile
import tempfile
import logging
import re
import fitz  # PyMuPDF
from docx import Document
from pathlib import Path
from typing import List
from fastapi import FastAPI, File, UploadFile, Form, HTTPException, BackgroundTasks
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from app.services.extraction import extract_literature_head, extract_metadata_via_llm, LiteratureMetadata

# ================= 1. 日志配置 =================
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger("DataClassifierEngine")

app = FastAPI(title="地质数据智能分类服务", version="1.0.0")
@app.get("/")
async def serve_frontend():
    """
    访问根路径时，直接返回前端的 index.html 页面
    """
    html_path = Path("frontend/index.html")
    if not html_path.exists():
        return {"error": "找不到前端文件，请检查 frontend/index.html 是否存在"}
    return FileResponse(html_path)

# 允许跨域
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ================= 2. 核心分类逻辑 =================
def extract_text_head_and_tail(file_path: str, max_pages=3) -> str:
    """
    掐头去尾提取法：提取前 max_pages 页和最后 1 页的文本。
    商业报告的免责声明通常在最后一页，文献的参考文献也在最后一页。
    """
    text = ""
    ext = Path(file_path).suffix.lower()
    try:
        if ext == '.pdf':
            with fitz.open(file_path) as doc:
                total_pages = len(doc)
                # 提取前几页
                for i in range(min(max_pages, total_pages)):
                    text += doc[i].get_text() + " "
                # 提取最后一页 (如果总页数大于头部提取页数)
                if total_pages > max_pages:
                    text += doc[total_pages - 1].get_text() + " "

        elif ext == '.docx':
            doc = Document(file_path)
            total_paras = len(doc.paragraphs)
            # 提取前几十段
            for i in range(min(max_pages * 10, total_paras)):
                text += doc.paragraphs[i].text + " "
            # 提取最后十段
            if total_paras > max_pages * 10:
                for i in range(total_paras - 10, total_paras):
                    text += doc.paragraphs[i].text + " "
    except Exception:
        pass

    text = text.lower()
    text = re.sub(r'\s+', ' ', text)
    return text.strip()


def classify_file_by_content(file_abspath: str) -> str:
    file_path = Path(file_abspath)
    ext = file_path.suffix.lower()
    name_lower = file_path.stem.lower()
    filename_display = file_path.name

    # 1. 明确的报表/演示格式 (如 .xlsx)
    if ext in {'.ppt', '.pptx', '.xls', '.xlsx'}:
        logger.info(f"[{filename_display}] 命中规则1: 明确报表后缀 -> 02-报告文档")
        return "02-报告文档"

    # 2. 过滤非文本类文件
    if ext not in {'.pdf', '.docx', '.doc', '.caj'}:
        return "99-其他待分类"

    # 3. 提取内容
    content = extract_text_head_and_tail(file_abspath)

    # 4. 加权特征词典匹配
    if content:
        # 结构：{词汇或正则: 权重分数}
        lit_dict = {
            r'doi:': 3, r'issn': 3, r'article info': 3, r'中图分类号': 3, r'基金项目': 3, r'handling editor': 3,
            r'abstract': 2, r'keywords': 2, r'elsevier': 2, r'received.*accepted': 2, r'references': 2,
            r'university': 2, r'摘要': 2, r'关键词': 2, r'参考文献': 2,
            r'introduction': 1, r'conclusions': 1, r'引言': 1
        }

        rep_dict = {
            r'all rights reserved': 3, r'disclaimer': 3, r'proprietary': 3, r'confidential': 3,
            r'commodity insights': 3, r'版权所有': 3, r'内部资料': 3,
            r'summary report': 2, r'table of contents': 2, r'contents': 2, r'prepared by': 2, r'client': 2,
            r'编制单位': 2, r'目录': 2,
            r'overview': 1, r'appendix': 1, r'附录': 1, r'概况': 1
        }

        lit_score = 0
        rep_score = 0
        hit_lit = []
        hit_rep = []

        # 遍历文献词典计分
        for pattern, weight in lit_dict.items():
            if re.search(pattern, content):
                lit_score += weight
                hit_lit.append(pattern)

        # 遍历报告词典计分
        for pattern, weight in rep_dict.items():
            if re.search(pattern, content):
                rep_score += weight
                hit_rep.append(pattern)

        if lit_score > 0 or rep_score > 0:
            if lit_score > rep_score:
                logger.info(
                    f"[{filename_display}] 评分机制 (文献 {lit_score} > 报告 {rep_score})。文献命中: {hit_lit}, 报告命中: {hit_rep} -> 01-公开文献")
                return "01-公开文献"
            else:
                logger.info(
                    f"[{filename_display}] 评分机制 (报告 {rep_score} >= 文献 {lit_score})。报告命中: {hit_rep}, 文献命中: {hit_lit} -> 02-报告文档")
                return "02-报告文档"

    # 5. 兜底策略：如果内容提取失败（图片PDF），使用文件名回退匹配
    report_name_keywords = ['report', 'summary', 'presentation', 'review', 'design', 'plan', '报告', '总结', '汇报',
                            '方案']
    if any(k in name_lower for k in report_name_keywords):
        logger.info(f"[{filename_display}] 内容无特征，文件名匹配 -> 02-报告文档")
        return "02-报告文档"

    return "01-公开文献" if ext == '.pdf' else "02-报告文档"




# ================= 3. 核心业务处理 =================
def process_classification(source_dir: Path, output_base_dir: str, target_area: str, mineral: str) -> str:
    """
    修改后：直接接收已准备好的 source_dir，不再负责解压。
    仅执行遍历、分类重构并重新打包的业务逻辑。
    """
    root_folder_name = f"{target_area}{mineral}"
    root_path = Path(output_base_dir) / root_folder_name
    root_path.mkdir(parents=True, exist_ok=True)

    categories = ["01-公开文献", "02-报告文档", "99-其他待分类"]
    for cat in categories:
        (root_path / cat).mkdir(exist_ok=True)

    processed_count = 0
    # 遍历 source_dir 进行分类
    for root, _, files in os.walk(source_dir):
        for file in files:
            source_file_path = Path(root) / file
            if file.startswith('.') or '__MACOSX' in str(source_file_path):
                continue

            category_folder = classify_file_by_content(str(source_file_path))
            target_file_path = root_path / category_folder / file

            counter = 1
            while target_file_path.exists():
                target_file_path = root_path / category_folder / f"{source_file_path.stem}_{counter}{source_file_path.suffix}"
                counter += 1

            shutil.copy2(source_file_path, target_file_path)
            processed_count += 1

    logger.info(f"Classification complete. Processed {processed_count} files.")

    output_zip_path = Path(output_base_dir) / f"{root_folder_name}_分类结果.zip"
    shutil.make_archive(str(output_zip_path).replace('.zip', ''), 'zip', root_dir=Path(output_base_dir),
                        base_dir=root_folder_name)

    return str(output_zip_path)

def cleanup_sandbox(temp_dir: str):
    """后台任务：在文件发送给用户后，立刻彻底删除服务器上的临时沙盒"""
    try:
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir)
            logger.info(f"沙盒回收成功，已释放服务器空间: {temp_dir}")
    except Exception as e:
        logger.error(f"沙盒回收失败: {temp_dir}, 错误: {str(e)}")
@app.post("/api/v1/classify")
async def upload_and_classify(
        background_tasks: BackgroundTasks,
        files: List[UploadFile] = File(...),  # 修改为接收文件列表
        target_area: str = Form(...),
        mineral: str = Form(...)
):
    if not files:
        raise HTTPException(status_code=400, detail="未接收到任何文件")

    temp_dir = tempfile.mkdtemp(prefix="data_class_")
    # 注册后台任务：当 FileResponse 成功将文件流发送给用户浏览器后，FastAPI 会自动在后台执行 cleanup_sandbox
    background_tasks.add_task(cleanup_sandbox, temp_dir)
    extract_dir = Path(temp_dir) / "extracted_source"
    extract_dir.mkdir(parents=True, exist_ok=True)

    try:
        # 场景 A: 上传的是单个 ZIP 压缩包
        if len(files) == 1 and files[0].filename.endswith('.zip'):
            zip_path = os.path.join(temp_dir, files[0].filename)
            with open(zip_path, "wb") as buffer:
                shutil.copyfileobj(files[0].file, buffer)
            try:
                with zipfile.ZipFile(zip_path, 'r') as zip_ref:
                    zip_ref.extractall(extract_dir)
            except zipfile.BadZipFile:
                raise ValueError("上传的文件不是有效的ZIP压缩包")

        # 场景 B: 上传的是文件夹（多个离散文件）
        else:
            for file in files:
                # 浏览器传文件夹时，filename 通常包含相对路径 (如 folder/sub/file.pdf)
                # 使用相对路径在沙盒中重建原始目录结构
                file_path = extract_dir / file.filename
                file_path.parent.mkdir(parents=True, exist_ok=True)
                with open(file_path, "wb") as buffer:
                    shutil.copyfileobj(file.file, buffer)

        logger.info(f"Received data for area {target_area} and mineral {mineral}")

        # 调用修改后的分类业务
        result_zip_path = process_classification(extract_dir, temp_dir, target_area, mineral)

        return FileResponse(
            path=result_zip_path,
            filename=os.path.basename(result_zip_path),
            media_type="application/zip"
        )

    except ValueError as ve:
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        logger.error(f"Internal Server Error: {str(e)}")
        raise HTTPException(status_code=500, detail="服务器内部处理错误，请联系管理员")


# ================= 5. 文献内容提取 API =================
@app.post("/api/v1/extract/literature", response_model=LiteratureMetadata)
async def auto_extract_literature_info(file: UploadFile = File(...)):
    """
    接收前端上传的单篇文献(PDF)，利用大模型自动解析并返回结构化元数据（用于前端表单预填）
    """
    if not file.filename.lower().endswith('.pdf'):
        raise HTTPException(status_code=400, detail="目前仅支持PDF格式文献的智能提取")

    # 创建临时文件保存上传的 PDF
    fd, temp_path = tempfile.mkstemp(suffix=".pdf")
    try:
        with os.fdopen(fd, 'wb') as f:
            shutil.copyfileobj(file.file, f)

        logger.info(f"开始智能解析文献: {file.filename}")

        # 1. 提取PDF文本 (调用 app.services.extraction 中的函数)
        doc_text = extract_literature_head(temp_path)
        if not doc_text:
            raise ValueError("未能提取到有效文字，可能是纯图片扫描件")

        # 2. 调用大模型提取元数据
        metadata_dict = extract_metadata_via_llm(doc_text)

        # 3. 日志记录并返回
        logger.info(f"文献 [{file.filename}] 解析成功: {metadata_dict}")
        return metadata_dict

    except ValueError as ve:
        logger.error(f"提取业务异常: {str(ve)}")
        raise HTTPException(status_code=400, detail=str(ve))
    except Exception as e:
        logger.error(f"提取系统异常: {str(e)}")
        raise HTTPException(status_code=500, detail="文献解析服务暂不可用，请稍后重试")
    finally:
        # 清理处理提取任务的单文件沙盒
        if os.path.exists(temp_path):
            os.remove(temp_path)